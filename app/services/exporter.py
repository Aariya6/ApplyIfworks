from pathlib import Path
import csv
from openpyxl import Workbook
from ..db import applications,source_health,companies
ROOT=Path(__file__).resolve().parents[2]
def export_tracker():
    rows=applications(); out=ROOT/"data/exports"; out.mkdir(parents=True,exist_ok=True)
    fields=["id","company","title","status","application_method","fit_score","application_friction","resume_version","resume_hash","submitted_resume_version","submitted_resume_hash","submitted_resume_filename","letter_version","created_at","updated_at","submitted_at","evidence_path","notes","apply_url","source"]
    wb=Workbook(); ws=wb.active; ws.title="Applications"; ws.append(fields)
    for r in rows: ws.append([r.get(f,"") for f in fields])
    ws2=wb.create_sheet("Source Health"); sh=source_health();
    if sh:
        keys=list(sh[0].keys()); ws2.append(keys)
        for r in sh: ws2.append([r.get(k,"") for k in keys])
    ws3=wb.create_sheet("Companies"); co=companies()
    if co:
        keys=list(co[0].keys()); ws3.append(keys)
        for r in co: ws3.append([r.get(k,"") for k in keys])
    wb.save(out/"applications.xlsx")
    with open(out/"applications.csv","w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows([{k:r.get(k,"") for k in fields} for r in rows])
    return str(out/"applications.xlsx"),str(out/"applications.csv")
