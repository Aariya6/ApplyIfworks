"""Auditable, local-first job intelligence pipeline.

The pipeline intentionally uses only sources whose terms and access have been reviewed.
It never submits an application.
"""

from __future__ import annotations

import csv
import datetime as dt
import difflib
import concurrent.futures
import collections
import hashlib
import html
import json
import logging
import re
import sqlite3
import threading
import time
import uuid
from dataclasses import asdict, dataclass, field
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import yaml
from openpyxl import Workbook

from ..adapters.router import run_source_report
from ..db import DB

ROOT = Path(__file__).resolve().parents[2]
LOGGER = logging.getLogger(__name__)
UTC = dt.timezone.utc
TRACKING_PARAMS = {"fbclid", "gclid", "ref", "source"}
ATS_HOSTS = (
    "boards.greenhouse.io",
    "job-boards.greenhouse.io",
    "jobs.lever.co",
    "jobs.ashbyhq.com",
    "jobs.smartrecruiters.com",
    "jobs.personio.com",
    "apply.workable.com",
)
DATE_KEYS = ("original_posted_at", "published_at", "publication_date", "datePublished", "releasedDate", "pubDate")


@dataclass(frozen=True)
class Evidence:
    snippet: str
    source_url: str


@dataclass(frozen=True)
class CriterionScore:
    points: float
    maximum: float
    evidence: tuple[Evidence, ...] = ()
    confidence: float = 0.0
    missing_evidence: tuple[str, ...] = ()
    reason_codes: tuple[str, ...] = ()


@dataclass(frozen=True)
class NormalizedJob:
    id: str
    requisition_id: str | None
    canonical_url: str
    source_url: str
    employer: str
    title: str
    location: str
    contract_type: str
    salary: str
    posted_at: str | None
    posted_date_source: str | None
    fetched_at: str
    source_id: str
    source_name: str
    remote_wording: str
    application_url: str
    description: str
    applicant_count: int | None
    raw: dict[str, Any] = field(repr=False)


@dataclass(frozen=True)
class JobAssessment:
    job: NormalizedJob
    eligibility: str
    eligibility_evidence: tuple[Evidence, ...]
    score: float
    criteria: dict[str, CriterionScore]
    hard_exclusions: tuple[str, ...]
    verification_reasons: tuple[str, ...]
    tier: str


def parse_date(value: Any) -> dt.datetime | None:
    """Parse common source timestamps without treating naive values as local time."""
    if value is None or value == "":
        return None
    try:
        if isinstance(value, (int, float)) or str(value).strip().isdigit():
            number = float(value)
            if number > 10_000_000_000:
                number /= 1000
            return dt.datetime.fromtimestamp(number, tz=UTC)
        text = str(value).strip()
        try:
            parsed = dt.datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            parsed = parsedate_to_datetime(text)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=UTC)
        return parsed.astimezone(UTC)
    except (OverflowError, OSError, TypeError, ValueError):
        return None


def canonicalize_url(url: str) -> str:
    if not url:
        return ""
    parsed = urlsplit(url.strip())
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        return ""
    query = [
        (key, value)
        for key, value in parse_qsl(parsed.query, keep_blank_values=True)
        if not key.lower().startswith("utm_") and key.lower() not in TRACKING_PARAMS
    ]
    path = parsed.path.rstrip("/") or "/"
    return urlunsplit((parsed.scheme.lower(), parsed.netloc.lower(), path, urlencode(query), ""))


def _raw_value(raw: dict[str, Any], names: Iterable[str]) -> Any:
    for name in names:
        value = raw.get(name)
        if value not in (None, ""):
            return value
    return None


def normalize_job(
    item: Any,
    source_name: str,
    source_url: str,
    fetched_at: dt.datetime | None = None,
) -> NormalizedJob:
    raw = getattr(item, "raw", None) or (item if isinstance(item, dict) else {})
    if not isinstance(raw, dict):
        raw = {}
    get = lambda key, default="": getattr(item, key, raw.get(key, default))
    date_keys = DATE_KEYS + (("date",) if source_name.casefold() == "remoteok" else ())
    posted_raw = _raw_value(raw, date_keys)
    if posted_raw is None and isinstance(item, dict):
        posted_raw = _raw_value(item, date_keys)
    posted = parse_date(posted_raw)
    app_url = str(
        get("apply_url")
        or _raw_value(raw, ("apply_url", "applyUrl", "hostedUrl", "absolute_url"))
        or ""
    )
    source_id = str(_raw_value(raw, ("id", "job_id", "posting_id")) or get("id") or "")
    req_id = _raw_value(raw, ("requisition_id", "requisitionId", "req_id", "job_requisition_id"))
    applicant_value = _raw_value(raw, ("applicant_count", "applicants_count", "number_of_applicants"))
    try:
        applicants = int(applicant_value) if applicant_value is not None else None
    except (TypeError, ValueError):
        applicants = None
    remote = " ".join(
        str(value or "")
        for value in (get("remote"), get("location"), raw.get("workplace_type"), raw.get("remote_policy"))
    ).strip()
    now = fetched_at or dt.datetime.now(UTC)
    url = str(get("discovery_url") or _raw_value(raw, ("source_url", "url", "share_link")) or app_url or "")
    return NormalizedJob(
        id=str(get("id") or hashlib.sha256(url.encode("utf-8")).hexdigest()[:20]),
        requisition_id=str(req_id) if req_id is not None else None,
        canonical_url=canonicalize_url(str(get("canonical_url") or app_url or url)),
        source_url=url,
        employer=str(get("company") or raw.get("company_name") or ""),
        title=str(get("title") or ""),
        location=str(get("location") or ""),
        contract_type=str(get("employment") or raw.get("employment_type") or ""),
        salary=str(get("salary") or raw.get("salary") or ""),
        posted_at=posted.isoformat() if posted else None,
        posted_date_source=next((key for key in date_keys if raw.get(key) not in (None, "")), None),
        fetched_at=now.astimezone(UTC).isoformat(),
        source_id=source_id,
        source_name=source_name,
        remote_wording=remote,
        application_url=app_url,
        description=str(get("description") or raw.get("description") or ""),
        applicant_count=applicants,
        raw=raw,
    )


def classify_eligibility(job: NormalizedJob) -> tuple[str, tuple[Evidence, ...], bool]:
    text = " ".join((job.remote_wording, job.location, job.description)).lower()
    remote = bool(re.search(r"\bfully remote\b|\bremote[- ]first\b|\bwork from anywhere\b", text))
    remote_claim = bool(re.search(r"\bremote\b", text))
    explicitly_nonremote = bool(
        re.search(r"\bhybrid\b|\bon[- ]site\b|\bin[- ]person\b|\bnot remote\b|\bno remote\b", text)
    )
    india_excluded = bool(
        re.search(
            r"\b(?:excluding|except|outside|not (?:open|available|eligible|permitted)|no hiring in)\b.{0,35}\bindia\b"
            r"|\bindia\b.{0,25}\bnot (?:supported|eligible|available|permitted|allowed)\b",
            text,
        )
    )
    india = bool(
        re.search(r"\bindia\b|\bworldwide\b|\bglobal(?:ly)?\b|\banywhere in the world\b|\ball countries\b", text)
    )
    contractor = bool(re.search(r"\bcontract(?:or|ors|ing)?\b|\bfreelance\b", text))
    contractor_excluded = bool(re.search(r"\bno contractors\b|\bcontractors? (?:are )?not (?:allowed|accepted)\b", text))
    restriction = bool(
        re.search(r"\b(only|must be based|based in|required to be located|residents? only)\b.{0,50}\b"
                  r"(us|united states|canada|europe|eu|uk|united kingdom|australia)\b", text)
    )
    if explicitly_nonremote:
        status = "not_eligible"
    elif india_excluded:
        status = "country_restricted"
    elif not remote:
        status = "unclear" if remote_claim else "not_eligible"
    elif restriction and not re.search(r"\bindia\b", text):
        status = "country_restricted"
    elif india and contractor and not contractor_excluded:
        status = "international_contractor_confirmed"
    elif india:
        status = "india_confirmed"
    else:
        status = "unclear"
    snippets: list[Evidence] = []
    for term in ("fully remote", "remote-first", "work from anywhere", "worldwide", "global", "India", "contractor", "must be based"):
        match = re.search(re.escape(term), f"{job.remote_wording} {job.location} {job.description}", re.I)
        if match:
            sentence = re.search(r"[^.!?\n]*" + re.escape(match.group(0)) + r"[^.!?\n]*", f"{job.remote_wording} {job.location} {job.description}", re.I)
            if sentence:
                snippets.append(Evidence(sentence.group(0).strip()[:240], job.source_url))
    return status, tuple(snippets), remote


def _is_authoritative_application_url(job: NormalizedJob) -> bool:
    if not job.application_url:
        return False
    parsed = urlsplit(job.application_url)
    if parsed.scheme.lower() != "https":
        return False
    host = parsed.hostname or ""
    if any(host == domain or host.endswith("." + domain) for domain in ATS_HOSTS):
        return True
    return False


def _profile_text(profile: dict[str, Any]) -> str:
    return json.dumps(profile, ensure_ascii=False, sort_keys=True).lower()


def _profile_skills(profile: dict[str, Any]) -> list[str]:
    skills = profile.get("skills", {})
    values = [value for group in skills.values() for value in group] if isinstance(skills, dict) else []
    return [str(value).strip() for value in values if str(value).strip()]


def _evidence_for_terms(text: str, terms: Iterable[str], source_url: str) -> list[Evidence]:
    found = []
    for term in terms:
        match = re.search(re.escape(term), text, re.I)
        if match:
            start = max(0, text.rfind(".", 0, match.start()) + 1)
            end = min(len(text), next((p for p in (text.find(".", match.end()), text.find("\n", match.end())) if p >= 0), len(text)))
            found.append(Evidence(text[start:end].strip()[:240], source_url))
    return found


def _criterion(
    points: float,
    maximum: float,
    evidence: list[Evidence],
    missing: str | Iterable[str],
    reason: str,
) -> CriterionScore:
    missing_items = (missing,) if isinstance(missing, str) else tuple(missing)
    confidence = min(1.0, points / maximum) if evidence and maximum else 0.0
    return CriterionScore(
        points=round(min(maximum, max(0.0, points)), 2),
        maximum=maximum,
        evidence=tuple(evidence),
        confidence=confidence,
        missing_evidence=missing_items,
        reason_codes=(reason,) if missing_items else (),
    )


def score_job(
    job: NormalizedJob,
    profile: dict[str, Any],
    now: dt.datetime | None = None,
) -> dict[str, CriterionScore]:
    text = f"{job.title}\n{job.description}"
    skills = _profile_skills(profile)
    matching = [skill for skill in skills if re.search(r"\b" + re.escape(skill) + r"\b", text, re.I)]
    skill_evidence = _evidence_for_terms(text, matching, job.source_url)
    skill_evidence.extend(Evidence(f"Profile lists skill: {skill}", "profile://config/profile.json") for skill in matching)
    missing_skills = [f"Profile skill not evidenced in posting: {skill}" for skill in skills if skill not in matching]
    skill_points = 25 * len(matching) / len(skills) if skills else 0

    eligibility, eligibility_evidence, _ = classify_eligibility(job)
    eligibility_points = 20 if eligibility in {"india_confirmed", "international_contractor_confirmed"} else 0
    eligibility_missing = [] if eligibility_points else ["Explicit India or international-contractor eligibility is missing."]

    experience = profile.get("experience", [])
    education = profile.get("education", [])
    seniority_terms = ("entry level", "junior", "graduate", "intern", "associate")
    seniority_evidence = _evidence_for_terms(text, seniority_terms, job.source_url)
    if seniority_evidence and (experience or education):
        seniority_points = 15
    else:
        seniority_points = 0
    if seniority_points and education:
        degree = education[0].get("degree", "Profile education")
        seniority_evidence.append(Evidence(f"Profile education: {degree}", "profile://config/profile.json"))
    seniority_missing = [] if seniority_points else ["Explicit role seniority or qualification-fit evidence is missing."]

    salary_evidence = _evidence_for_terms(job.description + " " + job.salary, (job.salary,) if job.salary else (), job.source_url)
    salary_points = 0
    if job.salary and re.search(r"\d", job.salary):
        salary_points = 10
    elif re.search(r"\b(paid|salary|compensation|hourly rate)\b", job.description, re.I):
        salary_points = 4
        salary_evidence = _evidence_for_terms(job.description, ("paid", "salary", "compensation", "hourly rate"), job.source_url)
    salary_missing = [] if job.salary else ["A compensation amount or range is missing."]

    posted = parse_date(job.posted_at)
    freshness_points = 0
    freshness_evidence: list[Evidence] = []
    if posted:
        age_hours = max(0.0, ((now or dt.datetime.now(UTC)).astimezone(UTC) - posted).total_seconds() / 3600)
        freshness_points = max(0.0, 10.0 * (1.0 - age_hours / (24.0 * 7.0)))
        freshness_evidence = [Evidence(f"Original publication date: {job.posted_at}", job.source_url)]
    if job.applicant_count is not None:
        freshness_points = min(10.0, freshness_points + 1.0)
        freshness_evidence.append(Evidence(f"Source explicitly reports {job.applicant_count} applicants.", job.source_url))
    freshness_missing = []
    if not posted:
        freshness_missing.append("Original publication date is missing.")
    if job.applicant_count is None:
        freshness_missing.append("Applicant count is unknown.")

    roles = profile.get("preferences", {}).get("preferred_domains", [])
    alignment_matches = [
        role for role in roles if re.search(re.escape(role), text, re.I)
    ]
    alignment_evidence = _evidence_for_terms(text, alignment_matches, job.source_url)
    alignment_evidence.extend(
        Evidence(f"Profile preference: {role}", "profile://config/profile.json")
        for role in alignment_matches
    )
    alignment_points = min(10.0, len(alignment_matches) * 3.0)
    alignment_missing = [
        f"Profile preference not evidenced in posting: {role}"
        for role in roles
        if role not in alignment_matches
    ]

    source_trust = 5 if _is_authoritative_application_url(job) else (2 if job.source_name.lower() in {"remoteok", "arbeitnow", "remotive"} else 0)
    trust_evidence = [Evidence(f"Application route uses recognized source {job.source_name}.", job.application_url)] if source_trust else []

    schedule_evidence = _evidence_for_terms(
        job.description,
        ("flexible hours", "overlap", "timezone", "utc", "working hours"),
        job.source_url,
    )
    schedule_points = min(5.0, len(schedule_evidence) * 2.5) if profile.get("preferences", {}).get("work_hours") else 0
    schedule_missing = []
    if not profile.get("preferences", {}).get("work_hours"):
        schedule_missing.append("Candidate work-hours preference is missing.")
    elif not schedule_evidence:
        schedule_missing.append("Posting schedule compatibility evidence is missing.")
    else:
        schedule_evidence.append(
            Evidence(
                f"Candidate work-hours preference: {profile['preferences']['work_hours']}",
                "profile://config/profile.json",
            )
        )
    if not profile.get("preferences", {}).get("work_hours"):
        schedule_evidence = []

    return {
        "documented_skills_match": _criterion(skill_points, 25, skill_evidence, missing_skills, "skills_evidence_missing"),
        "verified_india_eligibility": _criterion(eligibility_points, 20, list(eligibility_evidence), eligibility_missing, "eligibility_unverified"),
        "seniority_qualification_fit": _criterion(seniority_points, 15, seniority_evidence, seniority_missing, "seniority_fit_unverified"),
        "compensation_reliability": _criterion(salary_points, 10, salary_evidence, salary_missing, "compensation_unverified"),
        "freshness_competition": _criterion(freshness_points, 10, freshness_evidence, freshness_missing, "freshness_unverified"),
        "career_alignment_portability": _criterion(alignment_points, 10, alignment_evidence, alignment_missing, "career_alignment_unverified"),
        "source_employer_credibility": _criterion(source_trust, 5, trust_evidence, [] if source_trust else ["Authoritative source evidence is missing."], "source_trust_unverified"),
        "schedule_compatibility": _criterion(schedule_points, 5, schedule_evidence, schedule_missing, "schedule_unverified"),
    }


def _required_qualification_missing(job: NormalizedJob, profile: dict[str, Any]) -> bool:
    text = job.description.lower()
    profile_text = _profile_text(profile)
    for match in re.finditer(r"\b([a-z][a-z0-9+#.-]{1,30})\s+(?:is\s+)?required\b", text):
        qualification = match.group(1).lower()
        if qualification not in {"experience", "salary", "work", "remote", "location"} and qualification not in profile_text:
            return True
    return False


def assess_job(
    job: NormalizedJob,
    profile: dict[str, Any],
    *,
    since_hours: float,
    min_score: float,
    priority_score: float = 95,
    now: dt.datetime | None = None,
) -> JobAssessment:
    now = now or dt.datetime.now(UTC)
    eligibility, evidence, fully_remote = classify_eligibility(job)
    exclusions: list[str] = []
    verification: list[str] = []
    job_text = f"{job.remote_wording} {job.location} {job.description}"
    explicit_nonremote = bool(
        re.search(r"\bhybrid\b|\bon[- ]site\b|\bin[- ]person\b|\bnot remote\b|\bno remote\b", job_text, re.I)
    )
    if explicit_nonremote or not re.search(r"\bremote\b", job_text, re.I):
        exclusions.append("not_fully_remote")
    elif not fully_remote:
        verification.append("remote_scope_unclear")
    if eligibility == "country_restricted":
        exclusions.append("country_restricted")
    elif eligibility == "unclear":
        verification.append("india_eligibility_unclear")
    if re.search(r"\b(unpaid|uncompensated|volunteer only)\b", f"{job.title} {job.description}", re.I):
        exclusions.append("unpaid")
    if not job.salary:
        verification.append("compensation_unverified")
    if re.search(r"\b(closed|no longer accepting|position filled|applications are closed)\b", job.description, re.I):
        exclusions.append("closed")
    posted = parse_date(job.posted_at)
    if posted is None:
        verification.append("original_publication_date_missing")
    elif posted > now.astimezone(UTC) + dt.timedelta(minutes=5):
        verification.append("future_publication_timestamp")
    elif (now.astimezone(UTC) - posted).total_seconds() > since_hours * 3600:
        exclusions.append("stale")
    if _required_qualification_missing(job, profile):
        exclusions.append("required_qualification_missing")
    if re.search(r"\b(upfront fee|pay to apply|guaranteed job|wire transfer|crypto payment)\b", job.description, re.I):
        exclusions.append("scam_signal")
    if not _is_authoritative_application_url(job):
        exclusions.append("no_authoritative_application_route")
    if not profile.get("skills") and not profile.get("education") and not profile.get("experience"):
        verification.append("candidate_profile_has_no_verifiable_facts")

    criteria = score_job(job, profile, now=now)
    total = round(sum(value.points for value in criteria.values()), 2)
    if exclusions:
        tier = "excluded"
    elif verification:
        tier = "needs_verification"
    elif total >= max(95, priority_score):
        tier = "95_plus_priority"
    elif total >= max(90, min_score):
        tier = "90_plus_review"
    else:
        verification.append("below_review_threshold")
        tier = "needs_verification"
    return JobAssessment(
        job=job,
        eligibility=eligibility,
        eligibility_evidence=evidence,
        score=total,
        criteria=criteria,
        hard_exclusions=tuple(sorted(set(exclusions))),
        verification_reasons=tuple(sorted(set(verification))),
        tier=tier,
    )


def deduplicate_jobs(jobs: Iterable[NormalizedJob]) -> tuple[list[NormalizedJob], int]:
    kept: list[NormalizedJob] = []
    by_req: dict[str, int] = {}
    by_url: dict[str, int] = {}
    duplicates = 0
    for job in jobs:
        identity = None
        if job.requisition_id:
            identity = by_req.get(job.requisition_id.casefold())
        if identity is None and job.canonical_url:
            identity = by_url.get(job.canonical_url)
        if identity is None:
            for index, previous in enumerate(kept):
                if (
                    job.employer.casefold() == previous.employer.casefold()
                    and job.location.casefold() == previous.location.casefold()
                    and difflib.SequenceMatcher(None, job.title.casefold(), previous.title.casefold()).ratio() >= 0.92
                ):
                    identity = index
                    break
        if identity is not None:
            duplicates += 1
            if job.requisition_id:
                by_req[job.requisition_id.casefold()] = identity
            if job.canonical_url:
                by_url[job.canonical_url] = identity
            continue
        index = len(kept)
        kept.append(job)
        if job.requisition_id:
            by_req[job.requisition_id.casefold()] = index
        if job.canonical_url:
            by_url[job.canonical_url] = index
    return kept, duplicates


def claim_grounding(claims: Iterable[str], profile: dict[str, Any]) -> tuple[bool, list[str]]:
    source = _profile_text(profile)
    unsupported = []
    for claim in claims:
        value = re.sub(
            r"^(?:education|experience|project(?: technology)?|language):\s*",
            "",
            claim,
            flags=re.I,
        )
        parts = [part.strip().lower() for part in value.split("|") if part.strip()]
        if not parts or any(part not in source for part in parts):
            unsupported.append(claim)
    return not unsupported, unsupported


def validate_profile(profile: Any) -> dict[str, Any]:
    if not isinstance(profile, dict):
        raise ValueError("Candidate profile must be a JSON object.")
    expected = {
        "candidate": dict,
        "skills": dict,
        "education": list,
        "experience": list,
        "projects": list,
        "preferences": dict,
    }
    missing = [
        key for key, expected_type in expected.items()
        if not isinstance(profile.get(key), expected_type)
    ]
    if missing:
        raise ValueError("Candidate profile has missing or invalid sections: " + ", ".join(missing))
    if not str(profile["candidate"].get("name", "")).strip():
        raise ValueError("Candidate profile is missing candidate.name.")
    for section in ("education", "experience", "projects"):
        if any(not isinstance(item, dict) for item in profile[section]):
            raise ValueError(f"Candidate profile section {section} must contain objects.")
    if any(not isinstance(values, list) for values in profile["skills"].values()):
        raise ValueError("Candidate profile skills must map categories to lists.")
    return profile


def _load_registry(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as file:
        return list(csv.DictReader(file))


def _source_configs() -> dict[str, dict[str, Any]]:
    config = yaml.safe_load((ROOT / "config" / "sources.yaml").read_text(encoding="utf-8")) or {}
    sources = config.get("sources", {})
    result: dict[str, dict[str, Any]] = {}
    for name, cfg in sources.items():
        if not isinstance(cfg, dict):
            continue
        result[name.casefold()] = cfg
        for site in cfg.get("sites", []):
            if isinstance(site, dict):
                merged = {key: value for key, value in cfg.items() if key != "sites"}
                merged.update(site)
                merged["sites"] = [site]
                result.setdefault(str(site.get("name", name)).casefold(), merged)
    return result


def _terms_reviewed(value: str) -> bool:
    return value.strip().casefold() in {"true", "yes", "1", "reviewed"}


def _registry_source_config(entry: dict[str, str], configs: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    name = entry.get("source_name", "")
    configured = configs.get(name.casefold())
    url = entry.get("url", "")
    board_match = re.search(
        r"boards-api\.greenhouse\.io/v1/boards/([a-z0-9_-]+)/jobs",
        url,
        re.I,
    )
    if board_match:
        greenhouse = configs.get("greenhouse")
        if not greenhouse:
            return None
        return {
            **greenhouse,
            "adapter": "greenhouse",
            "companies": [board_match.group(1)],
            "access_method": entry.get("access_method", ""),
        }
    return configured


def _valid_sample_count(jobs: Iterable[Any]) -> int:
    valid = 0
    for item in jobs:
        title = getattr(item, "title", "") or (item.get("title", "") if isinstance(item, dict) else "")
        apply_url = getattr(item, "apply_url", "") or (item.get("apply_url", "") if isinstance(item, dict) else "")
        parsed = urlsplit(str(apply_url))
        if str(title).strip() and parsed.scheme.lower() == "https" and parsed.netloc:
            valid += 1
    return valid


def _safe_error_code(exc: Exception) -> str:
    response = getattr(exc, "response", None)
    status_code = getattr(response, "status_code", None)
    return f"HTTP_{status_code}" if status_code else type(exc).__name__


def _test_registry_source(
    entry: dict[str, str],
    configs: dict[str, dict[str, Any]],
    *,
    force: bool = False,
    before_attempt: Any = None,
) -> tuple[list[Any], str]:
    name = entry.get("source_name", "")
    if not _terms_reviewed(entry.get("terms_reviewed", "")):
        return [], "SKIPPED_TERMS_NOT_REVIEWED"
    if entry.get("enabled", "").strip().casefold() in {"false", "no", "0"}:
        return [], "SKIPPED_DISABLED"
    cfg = _registry_source_config(entry, configs)
    if not cfg or cfg.get("access_method", "").casefold() == "manual_import":
        return [], "MANUAL_OR_UNIMPLEMENTED"
    if str(cfg.get("enabled", True)).strip().casefold() in {"false", "no", "0"}:
        return [], "SKIPPED_DISABLED"
    if not force and entry.get("status", "").casefold() not in {"verified", "active"}:
        return [], "SKIPPED_NOT_VERIFIED"
    for attempt in range(3):
        if before_attempt:
            before_attempt()
        try:
            jobs, adapter_failures = run_source_report(name, cfg)
            valid_samples = _valid_sample_count(jobs)
            if not jobs:
                if adapter_failures:
                    if attempt < 2:
                        time.sleep(0.5 * (2**attempt))
                        continue
                    return [], f"FAIL:{adapter_failures[0]['error_type']}"
                return [], "ZERO_JOBS"
            if not valid_samples:
                if attempt < 2:
                    time.sleep(0.5 * (2**attempt))
                    continue
                return jobs, "INVALID_OUTPUT"
            if adapter_failures or valid_samples < len(jobs):
                if attempt < 2:
                    time.sleep(0.5 * (2**attempt))
                    continue
                issues = [failure["error_type"] for failure in adapter_failures]
                if valid_samples < len(jobs):
                    issues.append("INVALID_RECORDS")
                return jobs, "PARTIAL:" + ",".join(sorted(set(issues)))
            return jobs, "PASS"
        except Exception as exc:
            if attempt < 2:
                time.sleep(0.5 * (2**attempt))
                continue
            LOGGER.warning("Source failed: %s (%s)", name, type(exc).__name__)
            return [], f"FAIL:{_safe_error_code(exc)}"
    return [], "FAIL:UnknownError"


def _record_source_health(
    connection: sqlite3.Connection,
    source_name: str,
    status: str,
    jobs_found: int,
    valid_samples: int,
    checked_at: str,
) -> None:
    is_success = status == "PASS" or status.startswith("PARTIAL:")
    error_type = status.split(":", 1)[1] if ":" in status else (
        status if status in {"ZERO_JOBS", "INVALID_OUTPUT"} else None
    )
    connection.execute(
        "INSERT INTO pipeline_source_health(source_name,checked_at,status,jobs_found,error_type,valid_samples,last_success_at,last_failure_at) "
        "VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(source_name) DO UPDATE SET "
        "checked_at=excluded.checked_at,status=excluded.status,jobs_found=excluded.jobs_found,error_type=excluded.error_type,"
        "valid_samples=excluded.valid_samples,"
        "last_success_at=CASE WHEN excluded.last_success_at IS NOT NULL THEN excluded.last_success_at ELSE pipeline_source_health.last_success_at END,"
        "last_failure_at=CASE WHEN excluded.last_failure_at IS NOT NULL THEN excluded.last_failure_at ELSE pipeline_source_health.last_failure_at END",
        (
            source_name,
            checked_at,
            status,
            jobs_found,
            error_type,
            valid_samples,
            checked_at if is_success else None,
            checked_at if status.startswith(("FAIL:", "PARTIAL:")) or status in {"ZERO_JOBS", "INVALID_OUTPUT"} else None,
        ),
    )


def test_source(source_name: str, registry_path: Path | None = None) -> dict[str, Any]:
    registry_path = registry_path or ROOT / "config" / "pipeline_sources.csv"
    match = next(
        (entry for entry in _load_registry(registry_path) if entry["source_name"].casefold() == source_name.casefold()),
        None,
    )
    if match is None:
        raise ValueError(f"Unknown source: {source_name}")
    jobs, status = _test_registry_source(match, _source_configs(), force=True)
    checked_at = dt.datetime.now(UTC).isoformat()
    with sqlite3.connect(DB) as connection:
        _ensure_schema(connection)
        _record_source_health(connection, source_name, status, len(jobs), _valid_sample_count(jobs), checked_at)
    return {
        "source_name": source_name,
        "status": status,
        "jobs_found": len(jobs),
        "valid_samples": _valid_sample_count(jobs),
    }


def pipeline_health() -> list[dict[str, Any]]:
    with sqlite3.connect(DB) as connection:
        connection.row_factory = sqlite3.Row
        try:
            return [dict(row) for row in connection.execute(
                "SELECT source_name,checked_at,status,jobs_found,valid_samples,error_type,last_success_at,last_failure_at "
                "FROM pipeline_source_health ORDER BY source_name"
            )]
        except sqlite3.OperationalError as exc:
            if "no such table" in str(exc):
                return []
            raise


def pipeline_overview(registry_path: Path | None = None) -> dict[str, Any]:
    registry_path = registry_path or ROOT / "config" / "pipeline_sources.csv"
    with sqlite3.connect(DB, timeout=30) as connection:
        connection.row_factory = sqlite3.Row
        _ensure_schema(connection)
        run = connection.execute(
            "SELECT id,started_at,completed_at,summary_json FROM pipeline_runs "
            "ORDER BY started_at DESC LIMIT 1"
        ).fetchone()
        assessments: list[dict[str, Any]] = []
        summary: dict[str, Any] = {}
        if run:
            summary = json.loads(run["summary_json"] or "{}")
            for row in connection.execute(
                "SELECT assessment_json FROM pipeline_assessments WHERE run_id=? ORDER BY score DESC LIMIT 100",
                (run["id"],),
            ):
                assessments.append(json.loads(row["assessment_json"]))
        health = {
            row["source_name"]: dict(row)
            for row in connection.execute(
                "SELECT source_name,checked_at,status,jobs_found,valid_samples,error_type,last_success_at,last_failure_at "
                "FROM pipeline_source_health"
            )
        }
    sources = []
    for entry in _load_registry(registry_path):
        source_name = entry.get("source_name", "")
        sources.append(
            {
                **entry,
                "health": health.get(source_name),
            }
        )
    return {
        "run": dict(run) if run else None,
        "summary": summary,
        "assessments": assessments,
        "sources": sources,
    }


def _ensure_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS pipeline_runs(
          id TEXT PRIMARY KEY, started_at TEXT NOT NULL, completed_at TEXT,
          profile_version TEXT NOT NULL, parameters_json TEXT NOT NULL,
          summary_json TEXT NOT NULL DEFAULT '{}'
        );
        CREATE TABLE IF NOT EXISTS pipeline_jobs(
          job_id TEXT PRIMARY KEY, source_name TEXT NOT NULL, source_url TEXT NOT NULL,
          requisition_id TEXT, canonical_url TEXT, posted_at TEXT, posted_date_source TEXT,
          fetched_at TEXT NOT NULL, job_json TEXT NOT NULL,
          first_seen_at TEXT NOT NULL DEFAULT '', last_seen_at TEXT NOT NULL DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS pipeline_assessments(
          run_id TEXT NOT NULL, job_id TEXT NOT NULL, tier TEXT NOT NULL,
          score REAL NOT NULL, assessment_json TEXT NOT NULL,
          PRIMARY KEY(run_id, job_id), FOREIGN KEY(run_id) REFERENCES pipeline_runs(id)
        );
        CREATE TABLE IF NOT EXISTS pipeline_source_payloads(
          run_id TEXT NOT NULL, job_id TEXT NOT NULL, source_name TEXT NOT NULL,
          fetched_at TEXT NOT NULL, payload_json TEXT NOT NULL,
          PRIMARY KEY(run_id, job_id), FOREIGN KEY(run_id) REFERENCES pipeline_runs(id)
        );
        CREATE TABLE IF NOT EXISTS pipeline_tracker(
          job_id TEXT PRIMARY KEY, status TEXT NOT NULL, score REAL NOT NULL,
          original_url TEXT NOT NULL, posted_at TEXT, fetched_at TEXT NOT NULL,
          applied_at TEXT, resume_version TEXT, profile_version TEXT NOT NULL,
          letter_status TEXT NOT NULL, follow_up_date TEXT, outcome TEXT
        );
        CREATE TABLE IF NOT EXISTS pipeline_artifacts(
          id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT NOT NULL,
          profile_version TEXT NOT NULL, artifact_kind TEXT NOT NULL,
          artifact_path TEXT NOT NULL, sha256 TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS pipeline_source_health(
          source_name TEXT PRIMARY KEY, checked_at TEXT NOT NULL, status TEXT NOT NULL,
          jobs_found INTEGER NOT NULL, error_type TEXT, valid_samples INTEGER NOT NULL DEFAULT 0,
          last_success_at TEXT, last_failure_at TEXT
        );
        """
    )
    columns = {row[1] for row in connection.execute("PRAGMA table_info(pipeline_source_health)")}
    for name, declaration in (
        ("valid_samples", "INTEGER NOT NULL DEFAULT 0"),
        ("last_success_at", "TEXT"),
        ("last_failure_at", "TEXT"),
    ):
        if name not in columns:
            connection.execute(f"ALTER TABLE pipeline_source_health ADD COLUMN {name} {declaration}")
    job_columns = {row[1] for row in connection.execute("PRAGMA table_info(pipeline_jobs)")}
    for name in ("first_seen_at", "last_seen_at"):
        if name not in job_columns:
            connection.execute(f"ALTER TABLE pipeline_jobs ADD COLUMN {name} TEXT NOT NULL DEFAULT ''")


def _serialize_assessment(assessment: JobAssessment) -> dict[str, Any]:
    result = asdict(assessment)
    result["job"]["description"] = ""
    result["job"]["raw"] = {}
    return result


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "-", value).strip("-_")
    return cleaned[:60] or "unknown"


def _grounded_documents(
    assessment: JobAssessment,
    profile: dict[str, Any],
    profile_version: str,
    output_dir: Path,
) -> list[dict[str, str]]:
    job = assessment.job
    date_label = (parse_date(job.posted_at) or parse_date(job.fetched_at) or dt.datetime.now(UTC)).date().isoformat()
    stem = f"{_safe_name(job.employer)}-{_safe_name(job.title)}-{date_label}"
    facts: list[str] = []
    for education in profile.get("education", []):
        degree = education.get("degree")
        institution = education.get("institution")
        start = education.get("start")
        end = education.get("end")
        location = education.get("location")
        if degree or institution:
            facts.append("Education: " + " | ".join(filter(None, (degree, institution, start, end, location))))
    for experience in profile.get("experience", []):
        experience_facts = list(
            filter(
                None,
                (
                    experience.get("title"),
                    experience.get("company"),
                    experience.get("start"),
                    experience.get("end"),
                    experience.get("location"),
                ),
            )
        )
        if experience_facts:
            facts.append("Experience: " + " | ".join(experience_facts))
        facts.extend(str(bullet) for bullet in experience.get("bullets", []) if bullet)
    for project in profile.get("projects", []):
        project_name = str(project.get("name", "")).strip()
        if project_name:
            facts.append("Project: " + project_name)
        stack = project.get("stack", [])
        if isinstance(stack, (list, tuple)):
            facts.extend(f"Project technology: {technology}" for technology in stack if technology)
        elif stack:
            facts.append(f"Project technology: {stack}")
        facts.extend(str(bullet) for bullet in project.get("bullets", []) if bullet)
    languages = profile.get("candidate", {}).get("languages", {})
    if isinstance(languages, dict):
        facts.extend(f"Language: {language} | {level}" for language, level in languages.items())
    job_text = f"{job.title}\n{job.description}".casefold()
    facts.sort(key=lambda fact: sum(term.casefold() in job_text for term in _profile_skills(profile) if term), reverse=True)
    skills = [str(skill) for skill in _profile_skills(profile)]
    skills.sort(key=lambda skill: skill.casefold() not in job_text)
    candidate_name = str(profile.get("candidate", {}).get("name", "")).strip()
    candidate = profile.get("candidate", {})
    contact = [
        str(candidate.get(field, "")).strip()
        for field in ("email", "phone", "city", "linkedin", "github", "leetcode")
        if str(candidate.get(field, "")).strip()
    ]
    contact_line = " | ".join(contact)
    claims = ([candidate_name] if candidate_name else []) + ([contact_line] if contact_line else []) + facts + skills
    grounded, unsupported = claim_grounding(claims, profile)
    if not grounded:
        raise ValueError(f"Document claims are not grounded in profile: {len(unsupported)}")
    output_dir.mkdir(parents=True, exist_ok=True)
    candidate_name = candidate_name or "Candidate"
    resume_text = "\n".join(
        [
            candidate_name,
            *([contact_line] if contact_line else []),
            f"Application target: {job.title} at {job.employer}",
            "",
            "PROFILE FACTS",
            *facts,
            "",
            "SKILLS",
            ", ".join(skills),
        ]
    )
    letter_text = "\n".join(
        [
            f"Application draft for {job.title} at {job.employer}",
            "",
            "The statements below are copied from the candidate profile for human review.",
            *facts,
            "Relevant profile skills: " + ", ".join(skills),
        ]
    )
    generated = []
    for kind, content, extension in (
        ("resume", resume_text, "txt"),
        ("cover_letter", letter_text, "txt"),
    ):
        rendered = content + "\n"
        path = output_dir / f"{stem}-{kind}.{extension}"
        if path.exists() and path.read_text(encoding="utf-8") != rendered:
            path = output_dir / f"{stem}-{kind}-{hashlib.sha256(rendered.encode('utf-8')).hexdigest()[:8]}.{extension}"
        if not path.exists():
            path.write_text(rendered, encoding="utf-8")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        generated.append(
            {
                "job_id": job.id,
                "profile_version": profile_version,
                "resume_version": f"pipeline-{profile_version}",
                "kind": kind,
                "path": str(path),
                "sha256": digest,
                "grounding": "PASS",
            }
        )
    from docx import Document

    for kind, heading, content in (
        ("resume", "Resume draft", resume_text),
        ("cover_letter", "Cover letter draft", letter_text),
    ):
        digest_suffix = hashlib.sha256(content.encode("utf-8")).hexdigest()[:8]
        path = output_dir / f"{stem}-{kind}-{digest_suffix}.docx"
        if not path.exists():
            document = Document()
            document.add_heading(candidate_name, level=0)
            document.add_heading(heading, level=1)
            document.add_paragraph(f"Application target: {job.title} at {job.employer}")
            document.add_paragraph("Human review required; statements below are grounded in the candidate profile.")
            for line in content.splitlines():
                if line.strip():
                    document.add_paragraph(line)
            document.save(path)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        generated.append(
            {
                "job_id": job.id,
                "profile_version": profile_version,
                "resume_version": f"pipeline-{profile_version}",
                "kind": f"{kind}_docx",
                "path": str(path),
                "sha256": digest,
                "grounding": "PASS",
            }
        )
    return generated


def _export_report(assessments: list[JobAssessment], summary: dict[str, Any], run_id: str, export_dir: Path) -> None:
    export_dir.mkdir(parents=True, exist_ok=True)
    report = {
        "run_id": run_id,
        "summary": summary,
        "jobs": [_serialize_assessment(item) for item in assessments],
    }
    (export_dir / f"{run_id}.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    with (export_dir / f"{run_id}.csv").open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "job_id", "source_name", "source_id", "employer", "title", "tier", "score",
                "eligibility", "posted_at", "posted_date_source", "fetched_at", "original_url",
                "application_url", "reasons",
            ],
        )
        writer.writeheader()
        for item in assessments:
            writer.writerow(
                {
                    "job_id": item.job.id,
                    "source_name": item.job.source_name,
                    "source_id": item.job.source_id,
                    "employer": item.job.employer,
                    "title": item.job.title,
                    "tier": item.tier,
                    "score": item.score,
                    "eligibility": item.eligibility,
                    "posted_at": item.job.posted_at or "",
                    "posted_date_source": item.job.posted_date_source or "",
                    "fetched_at": item.job.fetched_at,
                    "original_url": item.job.source_url,
                    "application_url": item.job.application_url,
                    "reasons": ";".join(item.hard_exclusions + item.verification_reasons),
                }
            )
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Run Summary"
    summary_sheet.append(["Metric", "Value"])
    for key, value in summary.items():
        summary_sheet.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    jobs_sheet = workbook.create_sheet("Jobs")
    job_fields = [
        "job_id", "source_name", "source_id", "employer", "title", "tier", "score", "eligibility",
        "posted_at", "posted_date_source", "fetched_at", "original_url", "application_url", "reasons",
    ]
    jobs_sheet.append(job_fields)
    for item in assessments:
        jobs_sheet.append(
            [
                item.job.id, item.job.source_name, item.job.source_id, item.job.employer,
                item.job.title, item.tier, item.score, item.eligibility, item.job.posted_at or "",
                item.job.posted_date_source or "", item.job.fetched_at, item.job.source_url,
                item.job.application_url, ";".join(item.hard_exclusions + item.verification_reasons),
            ]
        )
    sources_sheet = workbook.create_sheet("Sources")
    sources_sheet.append(["Source", "Jobs found", "Status"])
    source_statuses = {
        row["source"]: row["reason"]
        for row in summary.get("source_skipped_details", [])
    }
    source_statuses.update(
        {row["source"]: row["error"] for row in summary.get("source_failures", [])}
    )
    for source_name, count in sorted(summary.get("source_counts", {}).items()):
        sources_sheet.append([source_name, count, source_statuses.get(source_name, "PASS")])
    workbook.save(export_dir / f"{run_id}.xlsx")
    report_rows = []
    for item in assessments:
        safe_source_url = canonicalize_url(item.job.source_url)
        source_cell = f"Source: {html.escape(item.job.source_name)}"
        if safe_source_url:
            source_cell += f" <a href=\"{html.escape(safe_source_url, quote=True)}\">listing</a>"
        safe_application_url = canonicalize_url(item.job.application_url)
        if safe_application_url:
            source_cell += f" <a href=\"{html.escape(safe_application_url, quote=True)}\">apply</a>"
        report_rows.append(
            f"<tr><td>{html.escape(item.tier)}</td><td>{item.score:.2f}</td>"
            f"<td>{html.escape(item.job.employer)}</td><td>{html.escape(item.job.title)}</td>"
            f"<td>{html.escape(item.eligibility)}</td><td>{source_cell}</td>"
            f"<td>{html.escape(item.job.posted_at or 'unknown')}</td>"
            f"<td>{html.escape(item.job.fetched_at)}</td></tr>"
        )
    rows = "".join(report_rows)
    (export_dir / f"{run_id}.html").write_text(
        "<!doctype html><meta charset=\"utf-8\"><title>Job pipeline report</title>"
        "<h1>Job pipeline report</h1><table><thead><tr><th>Tier</th><th>Score</th><th>Employer</th><th>Role</th><th>Eligibility</th><th>Source attribution and routes</th><th>Posted</th><th>Fetched</th></tr></thead>"
        f"<tbody>{rows}</tbody></table>",
        encoding="utf-8",
    )


def run_pipeline(
    *,
    since_hours: float = 24,
    min_score: float = 90,
    priority_score: float = 95,
    max_sources: int | None = None,
    max_jobs: int | None = None,
    concurrency: int = 1,
    requests_per_second: float = 1.0,
    refresh_source: str | None = None,
    dry_run: bool = True,
    export_dir: Path | None = None,
    application_output_dir: Path | None = None,
    registry_path: Path | None = None,
    profile_path: Path | None = None,
    now: dt.datetime | None = None,
) -> dict[str, Any]:
    if not dry_run:
        raise ValueError("This pipeline only prepares drafts; application submission is not implemented here.")
    if not 90 <= min_score <= 100:
        raise ValueError("min_score must be between 90 and 100; this pipeline does not lower its review floor.")
    if not 95 <= priority_score <= 100 or priority_score < min_score:
        raise ValueError("priority_score must be between min_score and 100, and cannot be below 95.")
    if concurrency < 1:
        raise ValueError("concurrency must be at least one.")
    if requests_per_second <= 0:
        raise ValueError("requests_per_second must be greater than zero.")
    if since_hours <= 0:
        raise ValueError("since_hours must be greater than zero.")
    if max_sources is not None and max_sources < 0:
        raise ValueError("max_sources cannot be negative.")
    if max_jobs is not None and max_jobs < 0:
        raise ValueError("max_jobs cannot be negative.")
    registry_path = registry_path or ROOT / "config" / "pipeline_sources.csv"
    profile_path = profile_path or ROOT / "config" / "profile.json"
    profile_bytes = profile_path.read_bytes()
    profile = validate_profile(json.loads(profile_bytes.decode("utf-8")))
    profile_version = hashlib.sha256(profile_bytes).hexdigest()[:16]
    started_at = dt.datetime.now(UTC).isoformat()
    run_id = dt.datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:8]
    configs = _source_configs()
    all_registry = _load_registry(registry_path)
    if not all_registry:
        raise ValueError(f"Source registry is empty: {registry_path}")
    if refresh_source:
        registry = [entry for entry in all_registry if entry["source_name"].casefold() == refresh_source.casefold()]
        if not registry:
            raise ValueError(f"Unknown source: {refresh_source}")
    else:
        registry = [
            entry for entry in all_registry
            if _terms_reviewed(entry.get("terms_reviewed", ""))
            and entry.get("status", "").casefold() in {"verified", "active"}
        ]
    if max_sources is not None:
        registry = registry[:max(0, max_sources)]
    if max_jobs == 0:
        registry = []

    parameters = {
        "since_hours": since_hours,
        "min_score": min_score,
        "priority_score": priority_score,
        "max_sources": max_sources,
        "max_jobs": max_jobs,
        "concurrency": concurrency,
        "requests_per_second": requests_per_second,
        "refresh_source": refresh_source,
        "dry_run": dry_run,
        "registry_path": str(registry_path),
        "profile_version": profile_version,
    }
    with sqlite3.connect(DB, timeout=30) as connection:
        _ensure_schema(connection)
        connection.execute(
            "INSERT INTO pipeline_runs(id,started_at,profile_version,parameters_json,summary_json) VALUES(?,?,?,?,?)",
            (run_id, started_at, profile_version, json.dumps(parameters), "{}"),
        )

    found: list[NormalizedJob] = []
    source_counts: dict[str, int] = {}
    source_failures: list[dict[str, str]] = []
    source_skipped: list[dict[str, str]] = []
    source_status_counts: collections.Counter[str] = collections.Counter(
        entry.get("status", "unknown").casefold() for entry in all_registry
    )
    successful_sources = 0
    partial_sources = 0
    limiter_lock = threading.Lock()
    next_request_at = [0.0]
    request_interval = 1.0 / max(0.1, requests_per_second)

    def wait_for_rate_limit() -> None:
        with limiter_lock:
            current = time.monotonic()
            wait = max(0.0, next_request_at[0] - current)
            next_request_at[0] = max(current, next_request_at[0]) + request_interval
        if wait:
            time.sleep(wait)

    def fetch_entry(entry: dict[str, str]) -> tuple[dict[str, str], list[Any], str]:
        name = entry.get("source_name", "")
        fetched, status = _test_registry_source(
            entry,
            configs,
            force=bool(refresh_source),
            before_attempt=wait_for_rate_limit,
        )
        return entry, fetched, status

    fetched_entries: list[tuple[dict[str, str], list[Any], str]] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, concurrency)) as executor:
        futures = {executor.submit(fetch_entry, entry): entry for entry in registry}
        for future in concurrent.futures.as_completed(futures):
            try:
                fetched_entries.append(future.result())
            except Exception as exc:
                entry = futures[future]
                fetched_entries.append((entry, [], f"FAIL:{_safe_error_code(exc)}"))
    fetched_entries.sort(key=lambda result: result[0].get("source_name", "").casefold())

    for entry, fetched, status in fetched_entries:
        source_name = entry.get("source_name", "")
        source_counts[source_name] = len(fetched)
        if status == "PASS":
            successful_sources += 1
        if status.startswith("PARTIAL:"):
            partial_sources += 1
        if status.startswith(("FAIL:", "PARTIAL:")) or status in {"ZERO_JOBS", "INVALID_OUTPUT"}:
            source_failures.append(
                {"source": source_name, "error": status.split(":", 1)[1] if ":" in status else status}
            )
        elif status != "PASS":
            source_skipped.append({"source": source_name, "reason": status})
        checked_at = dt.datetime.now(UTC).isoformat()
        for item in fetched:
            if max_jobs is not None and len(found) >= max_jobs:
                break
            found.append(normalize_job(item, source_name, entry.get("url", ""), dt.datetime.fromisoformat(checked_at)))
        with sqlite3.connect(DB, timeout=30) as connection:
            _ensure_schema(connection)
            _record_source_health(connection, source_name, status, len(fetched), _valid_sample_count(fetched), checked_at)

    deduped, duplicate_count = deduplicate_jobs(found)
    assessed = [
        assess_job(
            job,
            profile,
            since_hours=since_hours,
            min_score=min_score,
            priority_score=priority_score,
            now=now,
        )
        for job in deduped
    ]
    summary = {
        "sources_configured": len(all_registry),
        "sources_selected": len(registry),
        "sources_succeeded": successful_sources,
        "sources_partial": partial_sources,
        "sources_failed": len(source_failures) - partial_sources,
        "sources_skipped": max(0, len(all_registry) - len(registry)) + len(source_skipped),
        "source_registry_status_counts": dict(sorted(source_status_counts.items())),
        "source_failures": source_failures,
        "source_skipped_details": source_skipped,
        "source_counts": source_counts,
        "discovered": len(found),
        "deduplicated": duplicate_count,
        "stale": sum("stale" in item.hard_exclusions for item in assessed),
        "ineligible": sum(item.eligibility in {"country_restricted", "not_eligible"} for item in assessed),
        "unknown_eligibility": sum(item.eligibility == "unclear" for item in assessed),
        "qualified_by_tier": {
            tier: sum(item.tier == tier for item in assessed)
            for tier in ("90_plus_review", "95_plus_priority", "excluded", "needs_verification")
        },
        "documents_generated": 0,
        "errors": len(source_failures),
        "dry_run": dry_run,
        "profile_version": profile_version,
        "run_id": run_id,
    }
    connection = sqlite3.connect(DB, timeout=30)
    try:
        _ensure_schema(connection)
        for assessment in assessed:
            job = assessment.job
            serialized = _serialize_assessment(assessment)
            connection.execute(
                "INSERT INTO pipeline_jobs(job_id,source_name,source_url,requisition_id,canonical_url,posted_at,posted_date_source,fetched_at,job_json,first_seen_at,last_seen_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(job_id) DO UPDATE SET source_name=excluded.source_name,source_url=excluded.source_url,requisition_id=excluded.requisition_id,canonical_url=excluded.canonical_url,posted_at=excluded.posted_at,posted_date_source=excluded.posted_date_source,fetched_at=excluded.fetched_at,job_json=excluded.job_json,last_seen_at=excluded.last_seen_at,"
                "first_seen_at=CASE WHEN pipeline_jobs.first_seen_at='' THEN excluded.first_seen_at ELSE pipeline_jobs.first_seen_at END",
                (
                    job.id, job.source_name, job.source_url, job.requisition_id, job.canonical_url,
                    job.posted_at, job.posted_date_source, job.fetched_at,
                    json.dumps({**asdict(job), "raw": {}}, ensure_ascii=False), started_at, job.fetched_at,
                ),
            )
            connection.execute(
                "INSERT OR REPLACE INTO pipeline_assessments(run_id,job_id,tier,score,assessment_json) VALUES(?,?,?,?,?)",
                (run_id, job.id, assessment.tier, assessment.score, json.dumps(serialized, ensure_ascii=False)),
            )
            connection.execute(
                "INSERT INTO pipeline_source_payloads(run_id,job_id,source_name,fetched_at,payload_json) VALUES(?,?,?,?,?)",
                (run_id, job.id, job.source_name, job.fetched_at, json.dumps(job.raw, ensure_ascii=False)),
            )
            connection.execute(
                "INSERT INTO pipeline_tracker(job_id,status,score,original_url,posted_at,fetched_at,profile_version,letter_status) VALUES(?,?,?,?,?,?,?,?) "
                "ON CONFLICT(job_id) DO UPDATE SET score=excluded.score,original_url=excluded.original_url,posted_at=excluded.posted_at,fetched_at=excluded.fetched_at,profile_version=excluded.profile_version",
                (job.id, "DRAFT_REVIEW" if assessment.tier in {"90_plus_review", "95_plus_priority"} else assessment.tier.upper(), assessment.score, job.source_url, job.posted_at, job.fetched_at, profile_version, "NOT_GENERATED"),
            )
        connection.commit()
    finally:
        connection.close()

    export_dir = export_dir or ROOT / "data" / "exports"
    application_output_dir = application_output_dir or ROOT / "data" / "applications"
    manifest: list[dict[str, str]] = []
    document_errors: list[dict[str, str]] = []
    for assessment in assessed:
        if assessment.tier not in {"90_plus_review", "95_plus_priority"}:
            continue
        try:
            artifacts = _grounded_documents(
                assessment,
                profile,
                profile_version,
                application_output_dir / assessment.job.id / "pipeline" / profile_version,
            )
            manifest.extend(artifacts)
        except Exception as exc:
            LOGGER.error("Draft generation failed for job %s (%s)", assessment.job.id, type(exc).__name__)
            document_errors.append({"job_id": assessment.job.id, "error": type(exc).__name__})
    summary["documents_generated"] = len(manifest)
    summary["document_errors"] = document_errors
    summary["errors"] += len(document_errors)
    if manifest:
        manifest_path = export_dir / f"{run_id}-manifest.json"
        export_dir.mkdir(parents=True, exist_ok=True)
        manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        with sqlite3.connect(DB, timeout=30) as connection:
            _ensure_schema(connection)
            for artifact in manifest:
                connection.execute(
                    "INSERT INTO pipeline_artifacts(job_id,profile_version,artifact_kind,artifact_path,sha256,created_at) "
                    "SELECT ?,?,?,?,?,? WHERE NOT EXISTS (SELECT 1 FROM pipeline_artifacts "
                    "WHERE job_id=? AND profile_version=? AND artifact_kind=? AND sha256=?)",
                    (
                        artifact["job_id"], artifact["profile_version"], artifact["kind"], artifact["path"],
                        artifact["sha256"], dt.datetime.now(UTC).isoformat(), artifact["job_id"],
                        artifact["profile_version"], artifact["kind"], artifact["sha256"],
                    ),
                )
                connection.execute(
                    "UPDATE pipeline_tracker SET resume_version=?,letter_status=? WHERE job_id=?",
                    (
                        artifact["resume_version"],
                        "DRAFT_GENERATED" if artifact["kind"].startswith("cover_letter") else "NOT_GENERATED",
                        artifact["job_id"],
                    ),
                )
    _export_report(assessed, summary, run_id, export_dir)
    with sqlite3.connect(DB, timeout=30) as connection:
        connection.execute(
            "UPDATE pipeline_runs SET completed_at=?,summary_json=? WHERE id=?",
            (dt.datetime.now(UTC).isoformat(), json.dumps(summary), run_id),
        )
    return summary


TRACKER_STATUSES = {
    "DRAFT_REVIEW",
    "NEEDS_VERIFICATION",
    "EXCLUDED",
    "APPLIED",
    "WITHDRAWN",
    "REJECTED",
    "INTERVIEW",
    "OFFER",
}


def pipeline_tracker(status: str | None = None) -> list[dict[str, Any]]:
    with sqlite3.connect(DB, timeout=30) as connection:
        connection.row_factory = sqlite3.Row
        _ensure_schema(connection)
        query = (
            "SELECT job_id,status,score,original_url,posted_at,fetched_at,applied_at,resume_version,"
            "profile_version,letter_status,follow_up_date,outcome FROM pipeline_tracker"
        )
        values: tuple[str, ...] = ()
        if status:
            query += " WHERE status=?"
            values = (status.upper(),)
        query += " ORDER BY score DESC, fetched_at DESC"
        return [dict(row) for row in connection.execute(query, values)]


def update_pipeline_tracker(
    job_id: str,
    *,
    status: str | None = None,
    applied_at: str | None = None,
    follow_up_date: str | None = None,
    outcome: str | None = None,
) -> None:
    if status is not None and status.upper() not in TRACKER_STATUSES:
        raise ValueError(f"Unsupported tracker status: {status}")
    if status and status.upper() == "APPLIED" and not applied_at:
        raise ValueError("Set --applied-at from a confirmed application record before marking a job APPLIED.")
    if applied_at and not parse_date(applied_at):
        raise ValueError("applied_at must be a valid ISO-8601 timestamp or date.")
    if follow_up_date and not parse_date(follow_up_date):
        raise ValueError("follow_up_date must be a valid ISO-8601 timestamp or date.")
    updates = {
        "status": status.upper() if status else None,
        "applied_at": applied_at,
        "follow_up_date": follow_up_date,
        "outcome": outcome,
    }
    assignments = [f"{column}=?" for column, value in updates.items() if value is not None]
    values = tuple(value for value in updates.values() if value is not None)
    if not assignments:
        raise ValueError("Provide at least one tracker field to update.")
    with sqlite3.connect(DB, timeout=30) as connection:
        cursor = connection.execute(
            f"UPDATE pipeline_tracker SET {', '.join(assignments)} WHERE job_id=?",
            (*values, job_id),
        )
        if cursor.rowcount == 0:
            raise ValueError(f"Unknown pipeline tracker job ID: {job_id}")
