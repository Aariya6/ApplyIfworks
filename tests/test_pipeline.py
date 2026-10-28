import csv
import datetime as dt
import json
import sqlite3
from pathlib import Path
from types import SimpleNamespace
from openpyxl import load_workbook

from app.services import pipeline


UTC = dt.timezone.utc


def fixture_job():
    data = json.loads(Path("tests/fixtures/pipeline_job.json").read_text(encoding="utf-8"))
    return SimpleNamespace(**data)


def normalized_job(**updates):
    base = pipeline.normalize_job(fixture_job(), "greenhouse", "https://boards.greenhouse.io/")
    return pipeline.NormalizedJob(**{**base.__dict__, **updates})


def profile():
    return {
        "candidate": {"name": "Test Candidate", "is_india_based": True},
        "skills": {"Languages": ["Python"]},
        "education": [{"degree": "B.Sc Computer Science", "institution": "Example University"}],
        "experience": [],
        "projects": [],
        "preferences": {
            "preferred_domains": ["Python"],
            "work_hours": "UTC+5:30",
        },
    }


def test_parse_date_handles_epoch_iso_and_timezone():
    assert pipeline.parse_date("2026-10-28T02:00:00+02:00") == dt.datetime(2026, 10, 28, tzinfo=UTC)
    assert pipeline.parse_date("2026-10-28T00:00:00Z") == dt.datetime(2026, 10, 28, tzinfo=UTC)
    assert pipeline.parse_date(1793145600) == dt.datetime.fromtimestamp(1793145600, UTC)
    assert pipeline.parse_date("not a date") is None


def test_normalization_keeps_original_posting_and_fetch_times_separate():
    job = fixture_job()
    normalized = pipeline.normalize_job(
        job,
        "greenhouse",
        "https://boards.greenhouse.io/example",
        dt.datetime(2026, 10, 28, 1, tzinfo=UTC),
    )
    assert normalized.posted_at == "2026-10-28T00:00:00+00:00"
    assert normalized.posted_date_source == "published_at"
    assert normalized.fetched_at == "2026-10-28T01:00:00+00:00"
    assert normalized.requisition_id == "REQ-123"
    assert normalized.applicant_count == 4


def test_original_date_provenance_never_uses_aggregator_or_updated_timestamps():
    remoteok = SimpleNamespace(
        title="Engineer",
        company="Example",
        location="Remote",
        remote="Remote",
        apply_url="https://jobs.example.com/1",
        discovery_url="https://remoteok.com/1",
        raw={"date": "2026-10-28T00:00:00Z"},
    )
    assert pipeline.normalize_job(remoteok, "RemoteOK", "https://remoteok.com/api").posted_date_source == "date"
    for source_name, raw in (
        ("arbeitnow", {"created_at": "2026-10-28T00:00:00Z"}),
        ("greenhouse", {"updated_at": "2026-10-28T00:00:00Z"}),
    ):
        item = SimpleNamespace(
            title="Engineer",
            company="Example",
            location="Remote",
            remote="Remote",
            apply_url="https://jobs.example.com/1",
            discovery_url="https://example.com/1",
            raw=raw,
        )
        normalized = pipeline.normalize_job(item, source_name, "https://example.com")
        assert normalized.posted_at is None
        assert normalized.posted_date_source is None


def test_freshness_boundary_is_inclusive():
    job = normalized_job(posted_at="2026-10-28T00:00:00+00:00")
    now = dt.datetime(2026, 10, 29, tzinfo=UTC)
    exact = pipeline.assess_job(job, profile(), since_hours=24, min_score=90, now=now)
    late = pipeline.assess_job(
        pipeline.NormalizedJob(**{**job.__dict__, "posted_at": "2026-10-27T23:59:59+00:00"}),
        profile(),
        since_hours=24,
        min_score=90,
        now=now,
    )
    assert "stale" not in exact.hard_exclusions
    assert "stale" in late.hard_exclusions


def test_eligibility_requires_location_and_remote_evidence():
    job = normalized_job()
    assert pipeline.classify_eligibility(job)[0] == "international_contractor_confirmed"
    unclear = pipeline.NormalizedJob(
        **{
            **job.__dict__,
            "location": "Remote",
            "remote_wording": "Remote",
            "description": "Join our distributed engineering team.",
        }
    )
    status, _, fully_remote = pipeline.classify_eligibility(unclear)
    assert status == "unclear"
    assert not fully_remote
    assessment = pipeline.assess_job(
        unclear,
        profile(),
        since_hours=24,
        min_score=90,
        now=dt.datetime(2026, 10, 28, 1, tzinfo=UTC),
    )
    assert "remote_scope_unclear" in assessment.verification_reasons
    assert "not_fully_remote" not in assessment.hard_exclusions


def test_deduplication_uses_requisition_then_url_then_similarity():
    first = normalized_job()
    by_requisition = pipeline.NormalizedJob(**{**first.__dict__, "id": "second", "canonical_url": "https://other.example/job", "source_name": "other"})
    by_url = pipeline.NormalizedJob(**{**first.__dict__, "id": "third", "requisition_id": None})
    kept, duplicates = pipeline.deduplicate_jobs([first, by_requisition, by_url])
    assert kept == [first]
    assert duplicates == 2


def test_scoring_is_evidence_backed_and_hard_exclusions_override():
    job = normalized_job()
    criteria = pipeline.score_job(job, profile())
    assert sum(item.maximum for item in criteria.values()) == 100
    assert criteria["documented_skills_match"].points > 0
    assert criteria["documented_skills_match"].evidence[0].source_url == job.source_url
    assert criteria["schedule_compatibility"].points > 0

    excluded = pipeline.assess_job(
        pipeline.NormalizedJob(
            **{
                **job.__dict__,
                "description": "Hybrid in London, no remote work.",
                "location": "London",
                "remote_wording": "Hybrid",
            }
        ),
        profile(),
        since_hours=24,
        min_score=100,
        now=dt.datetime(2026, 10, 28, 1, tzinfo=UTC),
    )
    assert excluded.tier == "excluded"
    assert "not_fully_remote" in excluded.hard_exclusions


def test_hard_exclusions_cover_stale_unpaid_closed_missing_qualification_scam_and_route():
    base = normalized_job()
    now = dt.datetime(2026, 10, 28, 1, tzinfo=UTC)
    cases = [
        ({"posted_at": "2026-10-25T00:00:00+00:00"}, "stale", profile()),
        ({"description": base.description + " This is an unpaid role."}, "unpaid", profile()),
        ({"description": base.description + " The position is closed."}, "closed", profile()),
        ({"description": base.description + " Rust is required."}, "required_qualification_missing", profile()),
        ({"description": base.description + " Pay an upfront fee to apply."}, "scam_signal", profile()),
        ({"application_url": ""}, "no_authoritative_application_route", profile()),
    ]
    for changes, expected, candidate_profile in cases:
        assessment = pipeline.assess_job(
            pipeline.NormalizedJob(**{**base.__dict__, **changes}),
            candidate_profile,
            since_hours=24,
            min_score=90,
            now=now,
        )
        assert expected in assessment.hard_exclusions
        assert assessment.tier == "excluded"


def test_claim_grounding_flags_unsupported_claims():
    grounded, unsupported = pipeline.claim_grounding(["Python"], profile())
    assert grounded
    assert unsupported == []
    grounded, unsupported = pipeline.claim_grounding(["Led a 20-person team"], profile())
    assert not grounded
    assert unsupported == ["Led a 20-person team"]


def test_unknown_compensation_requires_verification_but_is_not_called_unpaid():
    base = normalized_job()
    job = pipeline.NormalizedJob(
        **{
            **base.__dict__,
            "salary": "",
            "description": "Fully remote worldwide including India. International contractors welcome. Python experience preferred.",
        }
    )
    assessment = pipeline.assess_job(
        job,
        profile(),
        since_hours=24,
        min_score=90,
        now=dt.datetime(2026, 10, 28, 1, tzinfo=UTC),
    )
    assert "compensation_unverified" in assessment.verification_reasons
    assert "unpaid" not in assessment.hard_exclusions


def test_source_test_requires_terms_review_and_records_health(tmp_path, monkeypatch):
    database = tmp_path / "pipeline.sqlite3"
    monkeypatch.setattr(pipeline, "DB", database)
    registry = tmp_path / "sources.csv"
    fields = ["source_name", "url", "source_type", "access_method", "terms_reviewed", "status"]
    with registry.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerow(
            {
                "source_name": "remoteok",
                "url": "https://remoteok.com/api",
                "source_type": "public_api",
                "access_method": "public_api",
                "terms_reviewed": "false",
                "status": "needs_review",
            }
        )
    monkeypatch.setattr(pipeline, "run_source_report", lambda *_args: (_ for _ in ()).throw(AssertionError("network call")))
    result = pipeline.test_source("remoteok", registry)
    assert result["status"] == "SKIPPED_TERMS_NOT_REVIEWED"
    assert pipeline.pipeline_health()[0]["status"] == "SKIPPED_TERMS_NOT_REVIEWED"

    entries = list(csv.DictReader(registry.open(encoding="utf-8")))
    entries[0]["terms_reviewed"] = "true"
    with registry.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(entries)
    monkeypatch.setattr(pipeline, "run_source_report", lambda *_args: ([], []))
    result = pipeline.test_source("remoteok", registry)
    assert result["status"] == "ZERO_JOBS"
    assert pipeline.pipeline_health()[0]["status"] == "ZERO_JOBS"


def test_source_test_gates_disabled_source_and_requires_valid_sample(tmp_path, monkeypatch):
    database = tmp_path / "pipeline.sqlite3"
    monkeypatch.setattr(pipeline, "DB", database)
    registry = tmp_path / "sources.csv"
    with registry.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=["source_name", "url", "source_type", "access_method", "terms_reviewed", "status", "enabled"])
        writer.writeheader()
        writer.writerow(
            {
                "source_name": "remoteok",
                "url": "https://remoteok.com/api",
                "source_type": "public_api",
                "access_method": "public_api",
                "terms_reviewed": "true",
                "status": "verified",
                "enabled": "false",
            }
        )
    monkeypatch.setattr(pipeline, "run_source_report", lambda *_args: (_ for _ in ()).throw(AssertionError("disabled source contacted")))
    assert pipeline.test_source("remoteok", registry)["status"] == "SKIPPED_DISABLED"

    entries = list(csv.DictReader(registry.open(encoding="utf-8")))
    entries[0]["enabled"] = "true"
    with registry.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=entries[0].keys())
        writer.writeheader()
        writer.writerows(entries)
    monkeypatch.setattr(
        pipeline,
        "run_source_report",
        lambda *_args: ([SimpleNamespace(title="Engineer", apply_url="https://jobs.example.com/1")], []),
    )
    result = pipeline.test_source("remoteok", registry)
    assert result["status"] == "PASS"
    assert result["valid_samples"] == 1


def test_source_config_disabled_gate_and_greenhouse_registry_routing(monkeypatch):
    entry = {
        "source_name": "remoteok",
        "url": "https://remoteok.com/api",
        "terms_reviewed": "true",
        "status": "verified",
    }
    monkeypatch.setattr(pipeline, "run_source_report", lambda *_args: (_ for _ in ()).throw(AssertionError("disabled adapter contacted")))
    jobs, status = pipeline._test_registry_source(
        entry,
        {"remoteok": {"adapter": "remoteok", "enabled": False}},
        force=True,
    )
    assert jobs == []
    assert status == "SKIPPED_DISABLED"

    greenhouse_entry = {
        "source_name": "GitLab Careers",
        "url": "https://boards-api.greenhouse.io/v1/boards/gitlab/jobs?content=true",
        "access_method": "official_public_api",
        "terms_reviewed": "true",
        "status": "verified",
    }
    calls = []

    def fake_report(name, config):
        calls.append((name, config))
        return [SimpleNamespace(title="Engineer", apply_url="https://boards.greenhouse.io/gitlab/jobs/123")], []

    monkeypatch.setattr(pipeline, "run_source_report", fake_report)
    jobs, status = pipeline._test_registry_source(
        greenhouse_entry,
        {"greenhouse": {"adapter": "greenhouse", "enabled": True, "companies": ["wrong"]}},
    )
    assert status == "PASS"
    assert len(jobs) == 1
    assert calls[0][0] == "GitLab Careers"
    assert calls[0][1]["companies"] == ["gitlab"]


def test_source_failure_retries_with_bounded_backoff(monkeypatch):
    attempts = []
    delays = []
    entry = {
        "source_name": "remoteok",
        "url": "https://remoteok.com/api",
        "terms_reviewed": "true",
        "status": "verified",
    }

    def flaky_source(*_args):
        attempts.append(True)
        if len(attempts) < 3:
            raise RuntimeError("temporary fixture failure")
        return [SimpleNamespace(title="Engineer", apply_url="https://jobs.example.com/1")], []

    monkeypatch.setattr(pipeline, "run_source_report", flaky_source)
    monkeypatch.setattr(pipeline.time, "sleep", delays.append)
    jobs, status = pipeline._test_registry_source(
        entry,
        {"remoteok": {"adapter": "remoteok", "enabled": True}},
    )
    assert status == "PASS"
    assert len(jobs) == 1
    assert len(attempts) == 3
    assert delays == [0.5, 1.0]


def test_pipeline_integration_uses_saved_fixture_and_exports_without_network(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "DB", tmp_path / "pipeline.sqlite3")
    profile_file = tmp_path / "profile.json"
    profile_file.write_text(json.dumps(profile()), encoding="utf-8")
    registry_file = tmp_path / "sources.csv"
    with registry_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=["source_name", "url", "source_type", "access_method", "terms_reviewed", "status"])
        writer.writeheader()
        writer.writerows(
            [
                {
                "source_name": "greenhouse",
                "url": "https://boards.greenhouse.io/",
                "source_type": "official_ats",
                "access_method": "official_api",
                "terms_reviewed": "true",
                "status": "verified",
                },
                {
                    "source_name": "remoteok",
                    "url": "https://remoteok.com/api",
                    "source_type": "public_api",
                    "access_method": "public_api",
                    "terms_reviewed": "true",
                    "status": "verified",
                },
            ]
        )
    saved_fixture = json.loads(Path("tests/fixtures/pipeline_job.json").read_text(encoding="utf-8"))
    def fake_source(name, _config):
        if name == "remoteok":
            raise RuntimeError("fixture failure")
        return [SimpleNamespace(**saved_fixture)], []

    monkeypatch.setattr(pipeline, "run_source_report", fake_source)
    monkeypatch.setattr(pipeline.time, "sleep", lambda _seconds: None)
    result = pipeline.run_pipeline(
        since_hours=24,
        min_score=90,
        priority_score=100,
        registry_path=registry_file,
        profile_path=profile_file,
        export_dir=tmp_path / "exports",
        application_output_dir=tmp_path / "applications",
        now=dt.datetime(2026, 10, 28, 1, tzinfo=UTC),
    )
    assert result["discovered"] == 1
    assert result["sources_succeeded"] == 1
    assert result["sources_failed"] == 1
    assert result["errors"] == 1
    assert result["qualified_by_tier"]["excluded"] == 0
    assert list((tmp_path / "exports").glob("*.json"))
    assert list((tmp_path / "exports").glob("*.csv"))
    assert list((tmp_path / "exports").glob("*.xlsx"))
    assert list((tmp_path / "exports").glob("*.html"))
    assert list((tmp_path / "exports").glob("*-manifest.json"))
    assert len(list((tmp_path / "applications").rglob("*.txt"))) == 2
    assert len(list((tmp_path / "applications").rglob("*.docx"))) == 2
    with sqlite3.connect(pipeline.DB) as connection:
        tracker = connection.execute(
            "SELECT status,posted_at,fetched_at,applied_at,resume_version,profile_version,letter_status "
            "FROM pipeline_tracker WHERE job_id='fixture-job-001'"
        ).fetchone()
    assert tracker[0] == "DRAFT_REVIEW"
    assert tracker[1] and tracker[2]
    assert tracker[3] is None
    assert tracker[4].startswith("pipeline-")
    assert tracker[5] and tracker[6] == "DRAFT_GENERATED"
    workbook = load_workbook(next((tmp_path / "exports").glob("*.xlsx")), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Run Summary", "Jobs", "Sources"]
    jobs_sheet = workbook["Jobs"]
    assert jobs_sheet.max_row == 2
    assert jobs_sheet["B2"].value == "greenhouse"
    assert jobs_sheet["L2"].value == "https://boards.greenhouse.io/example/jobs/123"
    workbook.close()
    overview = pipeline.pipeline_overview(registry_file)
    assert overview["run"]["id"] == result["run_id"]
    assert len(overview["assessments"]) == 1
    with sqlite3.connect(pipeline.DB) as connection:
        payload = connection.execute(
            "SELECT payload_json FROM pipeline_source_payloads WHERE job_id='fixture-job-001' ORDER BY rowid DESC LIMIT 1"
        ).fetchone()[0]
        assessment_json = connection.execute(
            "SELECT assessment_json FROM pipeline_assessments WHERE job_id='fixture-job-001' ORDER BY rowid DESC LIMIT 1"
        ).fetchone()[0]
    assert "Python experience preferred" in payload
    assert json.loads(assessment_json)["job"]["description"] == ""
    with sqlite3.connect(pipeline.DB) as connection:
        first_seen = connection.execute(
            "SELECT first_seen_at FROM pipeline_jobs WHERE job_id='fixture-job-001'"
        ).fetchone()[0]
        artifact_count = connection.execute(
            "SELECT COUNT(*) FROM pipeline_artifacts WHERE job_id='fixture-job-001'"
        ).fetchone()[0]

    second_result = pipeline.run_pipeline(
        since_hours=24,
        min_score=90,
        priority_score=100,
        registry_path=registry_file,
        profile_path=profile_file,
        export_dir=tmp_path / "exports",
        application_output_dir=tmp_path / "applications",
        now=dt.datetime(2026, 10, 28, 1, tzinfo=UTC),
    )
    assert second_result["documents_generated"] == 4
    assert pipeline.pipeline_overview(registry_file)["run"]["id"] == second_result["run_id"]
    assert len(list((tmp_path / "applications").rglob("*.txt"))) == 2
    assert len(list((tmp_path / "applications").rglob("*.docx"))) == 2
    with sqlite3.connect(pipeline.DB) as connection:
        assert connection.execute(
            "SELECT first_seen_at FROM pipeline_jobs WHERE job_id='fixture-job-001'"
        ).fetchone()[0] == first_seen
        assert connection.execute(
            "SELECT COUNT(*) FROM pipeline_artifacts WHERE job_id='fixture-job-001'"
        ).fetchone()[0] == artifact_count


def test_tracker_applied_status_requires_explicit_timestamp(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "DB", tmp_path / "tracker.sqlite3")
    with sqlite3.connect(pipeline.DB) as connection:
        pipeline._ensure_schema(connection)
        connection.execute(
            "INSERT INTO pipeline_tracker(job_id,status,score,original_url,fetched_at,profile_version,letter_status) "
            "VALUES('job-1','DRAFT_REVIEW',90,'https://example.com/job','2026-10-28T00:00:00+00:00','profile-v1','NOT_GENERATED')"
        )
    try:
        pipeline.update_pipeline_tracker("job-1", status="APPLIED")
    except ValueError as exc:
        assert "confirmed application record" in str(exc)
    else:
        raise AssertionError("Marking applied without timestamp should fail")
    pipeline.update_pipeline_tracker(
        "job-1",
        status="APPLIED",
        applied_at="2026-10-28T01:00:00Z",
        outcome="Submitted manually and confirmed by the candidate",
    )
    assert pipeline.pipeline_tracker()[0]["applied_at"] == "2026-10-28T01:00:00Z"


def test_source_registry_adds_30_review_only_user_candidates():
    rows = pipeline._load_registry(Path("config/pipeline_sources.csv"))
    names = {row["source_name"] for row in rows}
    urls = {row["url"] for row in rows}
    added = {
        "Dice", "Built In Jobs", "No Fluff Jobs", "Just Join IT", "Tech Jobs for Good",
        "Climatebase Jobs", "Work on Climate", "Impactpool", "ReliefWeb Jobs", "Idealist Jobs",
        "The Muse Jobs", "Poynter Jobs", "Journalism Jobs", "Mediabistro Jobs", "ProBlogger Jobs",
        "Write the Docs Jobs", "UX Writing Hub Jobs", "Kolabtree", "Nature Careers", "Jobs.ac.uk",
        "PhilJobs", "H-Net Jobs", "Game Jobs", "Hitmarker Jobs", "Work With Indies",
        "ArtStation Jobs", "Outlier AI", "DataAnnotation", "OneForma", "LXT Careers",
    }
    candidates = [row for row in rows if row["source_name"] in added]
    assert len(added) == 30
    assert len(candidates) == 30
    assert len(names) == len(rows)
    assert len(urls) == len(rows)
    assert all(
        row["status"] == "needs_review"
        and row["terms_reviewed"].lower() == "false"
        and row["access_method"] == "manual_import"
        for row in candidates
    )


def test_pipeline_and_sources_pages_show_real_registry_and_run_data(monkeypatch):
    source = {
        "source_name": "Dice",
        "source_type": "specialist_remote_board",
        "access_method": "manual_import",
        "terms_reviewed": "false",
        "status": "needs_review",
        "last_checked_at": "",
        "health": None,
    }
    job = {
        "source_url": "https://boards.greenhouse.io/example/jobs/123",
        "application_url": "https://boards.greenhouse.io/example/jobs/123",
        "employer": "Example Employer",
        "title": "Python Engineer",
        "posted_at": None,
        "source_name": "Example Careers",
    }
    overview = {
        "run": {
            "id": "fixture-run",
            "started_at": "2026-10-28T00:00:00+00:00",
            "completed_at": "2026-10-28T00:01:00+00:00",
        },
        "summary": {
            "sources_selected": 1,
            "sources_succeeded": 1,
            "discovered": 1,
            "deduplicated": 0,
            "documents_generated": 0,
        },
        "assessments": [{
            "tier": "needs_verification",
            "score": 0,
            "eligibility": "unclear",
            "job": job,
            "hard_exclusions": [],
            "verification_reasons": ["missing original publication date"],
        }],
        "sources": [source],
    }
    monkeypatch.setattr(pipeline, "pipeline_overview", lambda: overview)
    from app import main

    pipeline_response = main.pipeline_page()
    sources_response = main.sources()
    assert pipeline_response.status_code == 200
    assert sources_response.status_code == 200
    pipeline_html = pipeline_response.body.decode("utf-8")
    sources_html = sources_response.body.decode("utf-8")
    assert "fixture-run" in pipeline_html
    assert "Python Engineer" in pipeline_html
    assert "missing original publication date" in pipeline_html
    assert "boards.greenhouse.io/example/jobs/123" in pipeline_html
    assert "Run pipeline now (dry run)" in pipeline_html
    assert "Dice" in sources_html
    assert "manual_import" in sources_html
    assert "needs_review" in sources_html
    assert any(getattr(route, "path", None) == "/pipeline" for route in main.app.routes)
    assert any(getattr(route, "path", None) == "/pipeline/run" for route in main.app.routes)
    calls = {}

    def fake_run_pipeline(**kwargs):
        calls.update(kwargs)
        return {"run_id": "fixture-run"}

    monkeypatch.setattr(pipeline, "run_pipeline", fake_run_pipeline)
    run_response = main.pipeline_run_route()
    assert run_response.status_code == 303
    assert run_response.headers["location"] == "/pipeline?run_id=fixture-run"
    assert calls["dry_run"] is True
    assert calls["registry_path"].name == "pipeline_sources.csv"
    assert calls["profile_path"].name == "profile.json"
