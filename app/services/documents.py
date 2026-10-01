from pathlib import Path
import asyncio, html, os, shutil

def html_to_pdf(html_path:Path,pdf_path:Path):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        exe=os.getenv("RJP_CHROMIUM_EXECUTABLE") or shutil.which("chromium") or shutil.which("chromium-browser") or shutil.which("google-chrome")
        browser=p.chromium.launch(headless=True, executable_path=exe if exe else None)
        page=browser.new_page()
        page.set_content(html_path.read_text(encoding="utf-8"),wait_until="load")
        page.pdf(path=str(pdf_path),format="A4",print_background=True,margin={"top":"0","right":"0","bottom":"0","left":"0"})
        browser.close()

def make_docx(profile,job,path):
    try:
        from docx import Document
    except Exception:
        return
    d=Document(); c=profile["candidate"]; d.add_heading(c["name"],0); d.add_paragraph(f"{c['phone']} · {c['email']} · {c['city']} · {c['linkedin']}")
    d.add_heading("Professional Summary",1); d.add_paragraph("Computer Science undergraduate with hands-on experience in machine learning, cloud computing, AI software agents, full-stack development and technical leadership.")
    for heading,items in [("Education",profile["education"]),("Experience & Leadership",profile["experience"]),("Projects",profile["projects"])]:
        d.add_heading(heading,1)
        for x in items:
            name=x.get("institution") or x.get("title") or x.get("name"); d.add_heading(name,2); d.add_paragraph(str(x.get("company",x.get("degree",x.get("stack","")))))
            for b in x.get("bullets",[]): d.add_paragraph(b,style="List Bullet")
    d.add_heading("Technical Skills",1); d.add_paragraph("; ".join(f"{k}: {', '.join(v)}" for k,v in profile["skills"].items()))
    d.save(path)
